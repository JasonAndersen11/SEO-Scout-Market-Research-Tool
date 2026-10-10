import re
from datetime import datetime

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

NAVY = "1F2A44"
BLUE = "1F6FEB"
LIGHT = "F3F6FB"
GREEN = "E6F4EA"
THIN = Side(style="thin", color="D5DAE3")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _market_stats(city_cache: dict, niche: str, city: str, state: str):
    """Pull the Phase 2 volume/CPC for a market out of the city result cache."""
    first = niche.split()[0].lower() if niche else ""
    for key, text in city_cache.items():
        k_kw, k_city, k_state = (key.split("|") + ["", "", ""])[:3]
        if k_city == city.lower() and k_state.upper() == state.upper() and first in k_kw:
            vol = re.search(r"Monthly Volume:\s*(\d+)", text)
            cpc = re.search(r"CPC:\s*\$([\d.]+)", text)
            return (int(vol.group(1)) if vol else None, float(cpc.group(1)) if cpc else None)
    return (None, None)


def _sheet_name(base: str, used: set) -> str:
    name = re.sub(r"[\[\]\*\?/\\:]", "", base)[:31]
    candidate, n = name, 2
    while candidate.lower() in used:
        suffix = f" {n}"
        candidate = name[:31 - len(suffix)] + suffix
        n += 1
    used.add(candidate.lower())
    return candidate


def _is_ad_runner(row: dict) -> bool:
    return "ad" in row.get("Found Via", "").lower()


def build_workbook(leads: list[dict], city_cache: dict, parse) -> openpyxl.Workbook:
    """One Overview sheet plus one sheet per market, each with its prospects."""
    wb = openpyxl.Workbook()
    overview = wb.active
    overview.title = "Overview"
    used = {"overview"}

    markets = []
    for lead in sorted(leads, key=lambda l: (l["niche"], l["state"], l["city"])):
        rows = parse(lead["text"])
        if rows:
            markets.append((lead, rows))

    # ── Per-market sheets ───────────────────────────────────────────────
    sheet_names = []
    for lead, rows in markets:
        niche, city, state = lead["niche"].title(), lead["city"], lead["state"]
        name = _sheet_name(f"{niche} - {city} {state}", used)
        sheet_names.append(name)
        ws = wb.create_sheet(name)
        ws.sheet_view.showGridLines = False
        vol, cpc = _market_stats(city_cache, lead["niche"], city, state)

        ws.merge_cells("A1:I1")
        ws["A1"] = f"{niche} — {city}, {state}"
        ws["A1"].font = Font(size=16, bold=True, color="FFFFFF")
        ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
        ws["A1"].alignment = Alignment(vertical="center", indent=1)
        ws.row_dimensions[1].height = 30

        stats = [f"{len(rows)} prospects", f"{sum(_is_ad_runner(r) for r in rows)} running ads"]
        if vol is not None:
            stats.append(f"{vol}/mo searches")
        if cpc is not None:
            stats.append(f"${cpc:.2f} CPC")
        stats.append(f"Saved {lead['saved_at'].replace('T', ' ')}")
        ws.merge_cells("A2:I2")
        ws["A2"] = "   ·   ".join(stats)
        ws["A2"].font = Font(color="5B6577", italic=True)
        ws["A2"].alignment = Alignment(indent=1)
        ws.row_dimensions[2].height = 20

        headers = ["#", "Business", "Phone", "Website", "Found Via", "Owner", "Notes", "Called?", "Call Notes"]
        for col, h in enumerate(headers, 1):
            c = ws.cell(row=4, column=col, value=h)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor=BLUE)
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = BORDER
        ws.row_dimensions[4].height = 22

        for i, r in enumerate(rows, 1):
            row_idx = 4 + i
            site = r["Website"].strip()
            values = [i, r["Business Name"], r["Phone"], site, r["Found Via"], r["Owner"], r["Notes"], "", ""]
            fill = PatternFill("solid", fgColor=GREEN if _is_ad_runner(r) else (LIGHT if i % 2 == 0 else "FFFFFF"))
            for col, val in enumerate(values, 1):
                c = ws.cell(row=row_idx, column=col, value=val)
                c.fill = fill
                c.border = BORDER
                c.alignment = Alignment(wrap_text=True, vertical="top",
                                        horizontal="center" if col in (1, 8) else "left")
            ws.cell(row=row_idx, column=2).font = Font(bold=True)
            if site and not re.match(r"^(n/?a|unknown|no website)$", site, re.I):
                href = site if re.match(r"^https?://", site) else "https://" + site
                ws.cell(row=row_idx, column=4).hyperlink = href
                ws.cell(row=row_idx, column=4).font = Font(color=BLUE, underline="single")

        dv = DataValidation(type="list", formula1='"Yes,No,Voicemail,Callback,Not interested"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"H5:H{4 + len(rows)}")

        for col, w in enumerate([5, 30, 17, 32, 20, 18, 42, 16, 36], 1):
            ws.column_dimensions[get_column_letter(col)].width = w
        ws.freeze_panes = "A5"
        ws.auto_filter.ref = f"A4:I{4 + len(rows)}"
        ws.sheet_properties.tabColor = BLUE

    # ── Overview sheet ──────────────────────────────────────────────────
    overview.sheet_view.showGridLines = False
    overview.merge_cells("A1:G1")
    overview["A1"] = "Rank & Rent — Leads by Market"
    overview["A1"].font = Font(size=18, bold=True, color="FFFFFF")
    overview["A1"].fill = PatternFill("solid", fgColor=NAVY)
    overview["A1"].alignment = Alignment(vertical="center", indent=1)
    overview.row_dimensions[1].height = 34
    overview.merge_cells("A2:G2")
    overview["A2"] = f"{len(markets)} markets · {sum(len(r) for _, r in markets)} prospects · updated {datetime.now():%Y-%m-%d %H:%M}"
    overview["A2"].font = Font(color="5B6577", italic=True)
    overview["A2"].alignment = Alignment(indent=1)

    heads = ["Niche", "City", "State", "Prospects", "Running ads", "Searches/mo", "CPC"]
    for col, h in enumerate(heads, 1):
        c = overview.cell(row=4, column=col, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=BLUE)
        c.alignment = Alignment(horizontal="center")
        c.border = BORDER
    for i, ((lead, rows), name) in enumerate(zip(markets, sheet_names), 1):
        vol, cpc = _market_stats(city_cache, lead["niche"], lead["city"], lead["state"])
        vals = [lead["niche"].title(), lead["city"], lead["state"], len(rows),
                sum(_is_ad_runner(r) for r in rows), vol, cpc]
        for col, val in enumerate(vals, 1):
            c = overview.cell(row=4 + i, column=col, value=val)
            c.border = BORDER
            c.fill = PatternFill("solid", fgColor=LIGHT if i % 2 == 0 else "FFFFFF")
            c.alignment = Alignment(horizontal="left" if col <= 2 else "center")
        link = overview.cell(row=4 + i, column=2)
        link.hyperlink = f"#'{name}'!A1"
        link.font = Font(color=BLUE, underline="single", bold=True)
        if cpc is not None:
            overview.cell(row=4 + i, column=7).number_format = '"$"0.00'
    if not markets:
        overview["A5"] = "No leads saved yet. Run research for a city to fill this in."
        overview["A5"].font = Font(italic=True, color="5B6577")
    for col, w in enumerate([18, 22, 8, 12, 14, 14, 10], 1):
        overview.column_dimensions[get_column_letter(col)].width = w
    overview.freeze_panes = "A5"
    return wb
